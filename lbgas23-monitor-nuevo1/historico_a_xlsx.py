#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
historico.csv -> historico.xlsx
===============================

Convierte la serie de promedios diarios en un libro de Excel con tres hojas:

    Promedios   serie completa, con filtros y encabezado fijo
    Resumen     último corte por producto y ámbito nacional
    Metadatos   fuente, fecha de corte y cobertura

Se ejecuta dentro del flujo diario, después de xml_a_csv.py. El tablero no
carga ninguna librería para esto: el botón "Histórico" descarga el archivo ya
generado.

    python3 historico_a_xlsx.py historico.csv historico.xlsx
"""

import csv
import io
import os
import sys
from datetime import datetime, timezone

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: pip install openpyxl")

MARINO = "1E3A5F"
CLARO = "F0F4F8"


def encabezar(hoja, columnas):
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF", size=11)
        celda.fill = PatternFill("solid", fgColor=MARINO)
        celda.alignment = Alignment(horizontal="center", vertical="center")
    hoja.freeze_panes = "A2"


def ajustar(hoja, anchos):
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho


def main():
    origen = sys.argv[1] if len(sys.argv) > 1 else "historico.csv"
    destino = sys.argv[2] if len(sys.argv) > 2 else "historico.xlsx"

    if not os.path.exists(origen):
        sys.exit("No se encontró %s" % origen)

    with io.open(origen, encoding="utf-8-sig", newline="") as fh:
        filas = list(csv.DictReader(fh))
    if not filas:
        sys.exit("%s está vacío" % origen)

    fechas = sorted({(f.get("Fecha") or "").strip() for f in filas if f.get("Fecha")})
    ultima = fechas[-1]

    libro = Workbook()

    # ---------------------------------------------------- Promedios ------
    h = libro.active
    h.title = "Promedios"
    encabezar(h, ["Fecha", "Ámbito", "Clave", "Producto", "Promedio", "Estaciones"])
    for f in filas:
        try:
            promedio = round(float(f["Promedio"]), 4)
        except (ValueError, TypeError, KeyError):
            promedio = None
        try:
            estaciones = int(f["Estaciones"])
        except (ValueError, TypeError, KeyError):
            estaciones = None
        h.append([f.get("Fecha", ""), f.get("Ambito", ""), f.get("Clave", ""),
                  f.get("Producto", ""), promedio, estaciones])
    for fila in h.iter_rows(min_row=2, min_col=5, max_col=5):
        for celda in fila:
            celda.number_format = '"$"#,##0.00'
    h.auto_filter.ref = h.dimensions
    ajustar(h, [12, 12, 30, 12, 12, 12])

    # ------------------------------------------------------ Resumen ------
    r = libro.create_sheet("Resumen")
    encabezar(r, ["Producto", "Promedio nacional", "Estaciones", "Fecha del corte"])
    for f in filas:
        if f.get("Fecha") == ultima and f.get("Ambito") == "nacional":
            try:
                r.append([f.get("Producto", ""), round(float(f["Promedio"]), 2),
                          int(f["Estaciones"]), ultima])
            except (ValueError, TypeError):
                continue
    for fila in r.iter_rows(min_row=2, min_col=2, max_col=2):
        for celda in fila:
            celda.number_format = '"$"#,##0.00'
    ajustar(r, [14, 20, 14, 18])

    # ---------------------------------------------------- Metadatos ------
    m = libro.create_sheet("Metadatos")
    encabezar(m, ["Concepto", "Valor"])
    estados = len({f["Clave"] for f in filas if f.get("Ambito") == "estado"})
    datos = [
        ["Fuente", "Comisión Nacional de Energía (CNE)"],
        ["Leyenda oficial", "Valores estimados por la SENER con información de la CNE y el SAT"],
        ["Primer corte", fechas[0]],
        ["Último corte", ultima],
        ["Días en la serie", len(fechas)],
        ["Registros", len(filas)],
        ["Ámbitos", "nacional, estado, región"],
        ["Entidades con dato", estados],
        ["Generado", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")],
        ["Plataforma", "Monitor de Precios LB GAS 23 · Servicio Bautista"],
    ]
    for d in datos:
        m.append(d)
    for fila in m.iter_rows(min_row=2, max_col=1):
        for celda in fila:
            celda.font = Font(bold=True)
            celda.fill = PatternFill("solid", fgColor=CLARO)
    ajustar(m, [24, 64])

    libro.save(destino)
    print("Generado: %s (%d registros · %s a %s)"
          % (destino, len(filas), fechas[0], ultima))


if __name__ == "__main__":
    main()
